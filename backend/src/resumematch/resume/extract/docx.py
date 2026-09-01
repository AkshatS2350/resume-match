"""Deterministic local DOCX extraction without relationship traversal."""

from pathlib import Path
from typing import Literal
from zipfile import ZipFile

from docx import Document

from resumematch.core.schemas.extracted_text import ExtractedBlock, ExtractedText
from resumematch.resume.extract.hardening import parse_docx_xml
from resumematch.resume.extract.normalize import normalize


def extract_docx(path: Path) -> ExtractedText:
    with ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            if name.endswith(".xml"):
                parse_docx_xml(archive.read(name))
    document = Document(str(path))
    items: list[tuple[Literal["paragraph", "table_cell"], str]] = [
        ("paragraph", paragraph.text) for paragraph in document.paragraphs
    ]
    items.extend(
        ("table_cell", cell.text)
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )
    blocks: list[ExtractedBlock] = []
    text_parts: list[str] = []
    for ordinal, (layout_kind, raw) in enumerate(items, start=1):
        text = normalize(raw)
        if not text:
            continue
        start = sum(len(item) + 1 for item in text_parts)
        text_parts.append(text)
        blocks.append(
            ExtractedBlock(
                block_id=f"1:{ordinal}",
                section_id="preamble",
                page=1,
                start_offset=start,
                end_offset=start + len(text),
                text=text,
                layout_kind=layout_kind,
                column_index=None,
            )
        )
    return ExtractedText(
        text="\n".join(text_parts),
        blocks=tuple(blocks),
        page_count=1,
        pages_with_text_layer=(1,),
        extractor_version="python-docx@1.1.2",
    )
